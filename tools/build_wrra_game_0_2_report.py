#!/usr/bin/env python3
"""Build the WRRA-Game 0.2 Korean research report."""

from __future__ import annotations

import hashlib
import json
import math
from pathlib import Path

import matplotlib.pyplot as plt
from matplotlib import font_manager
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "wrra_game_output" / "v0_2"
RESULTS_PATH = ROOT / "wrra_game_output" / "wrra_game_0_2_results.json"
CODE_PATH = ROOT / "wrra_game_0_2.py"
DOCX_PATH = ROOT / "WRRA_Game_0.2_완전탐색_검증_오목과_바둑_최원식_2026-09-17.docx"
FONT = "Noto Sans KR"
MATH_FONT = "DejaVu Sans"
NAVY = "24364B"
PALE = "F3F6F8"


def korean_font() -> str:
    names = {f.name for f in font_manager.fontManager.ttflist}
    for candidate in ("WRRA Korean Sans", "Noto Sans KR", "Noto Sans CJK KR", "DejaVu Sans"):
        if candidate in names:
            return candidate
    return "DejaVu Sans"


plt.rcParams["font.family"] = korean_font()
plt.rcParams["axes.unicode_minus"] = False


def set_font(run, size=None, bold=None, color=None, italic=None, family=FONT):
    run.font.name = family
    rpr = run._element.get_or_add_rPr()
    for slot in ("eastAsia", "ascii", "hAnsi"):
        rpr.rFonts.set(qn(f"w:{slot}"), family)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:cantSplit")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def set_cell_width(cell, inches: float):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.8)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    normal.paragraph_format.line_spacing = 1.2
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True

    title = doc.styles["Title"]
    title.font.name = FONT
    title._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(12)
    ppr = title.element.get_or_add_pPr()
    border = ppr.find(qn("w:pBdr"))
    if border is not None:
        ppr.remove(border)

    for name, size, before, after in (
        ("Heading 1", 16, 17, 7),
        ("Heading 2", 13, 12, 5),
        ("Heading 3", 11.5, 9, 4),
    ):
        style = doc.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    cap = doc.styles["Caption"]
    cap.font.name = FONT
    cap._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    cap.font.size = Pt(9)
    cap.font.color.rgb = RGBColor(55, 55, 55)
    cap.paragraph_format.space_before = Pt(3)
    cap.paragraph_format.space_after = Pt(9)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    start = OxmlElement("w:fldChar")
    start.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((start, instr, end))
    set_font(run, 8.5, color="666666")


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def paragraph(doc, text="", bold_lead=None, align=None, size=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    if bold_lead and text.startswith(bold_lead):
        set_font(p.add_run(bold_lead), size=size, bold=True)
        set_font(p.add_run(text[len(bold_lead):]), size=size)
    else:
        set_font(p.add_run(text), size=size)
    return p


def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        set_font(p.add_run(item))


def numbered(doc, items):
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.23)
        p.paragraph_format.first_line_indent = Inches(-0.23)
        p.paragraph_format.space_after = Pt(3)
        set_font(p.add_run(f"{index}.  "))
        set_font(p.add_run(item))


def table(doc, headers, rows, widths, center_cols=(0,)):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    set_table_borders(tbl)
    hdr = tbl.rows[0]
    repeat_header(hdr)
    prevent_row_split(hdr)
    for idx, label in enumerate(headers):
        cell = hdr.cells[idx]
        set_cell_width(cell, widths[idx])
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(label), 9.0, True, "FFFFFF")
    for ridx, row in enumerate(rows):
        cells = tbl.add_row().cells
        prevent_row_split(tbl.rows[-1])
        for idx, value in enumerate(row):
            cell = cells[idx]
            set_cell_width(cell, widths[idx])
            if ridx % 2:
                set_cell_shading(cell, PALE)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx in center_cols else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            set_font(p.add_run(str(value)), 9.0)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)
    return tbl


def figure(doc, path: Path, caption: str, width: float):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(width))
    cp = doc.add_paragraph(style="Caption")
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(cp.add_run(caption), 9)


def render_equation(tex: str, filename: str, width=10.0, height=0.75, fontsize=19) -> Path:
    path = OUT / filename
    fig = plt.figure(figsize=(width, height), facecolor="white")
    fig.text(0.5, 0.5, f"${tex}$", ha="center", va="center", fontsize=fontsize, fontfamily=MATH_FONT)
    fig.savefig(path, dpi=260, bbox_inches="tight", pad_inches=0.08, facecolor="white")
    plt.close(fig)
    return path


def equation(doc, path: Path, width=5.9):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(7)
    p.add_run().add_picture(str(path), width=Inches(width))


def draw_go_board(ax, rows, title, highlight=None):
    n = len(rows)
    ax.set_facecolor("#D8AA66")
    for i in range(n):
        ax.plot([0, n - 1], [i, i], color="#3A2A1F", lw=1.2)
        ax.plot([i, i], [0, n - 1], color="#3A2A1F", lw=1.2)
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            if value == ".":
                continue
            color = "#111111" if value == "B" else "#F5F5F5"
            edge = "#000000" if value == "B" else "#777777"
            ax.scatter(c, n - 1 - r, s=720, color=color, edgecolors=edge, linewidths=1.2, zorder=3)
    if highlight is not None:
        r, c = highlight
        ax.scatter(c, n - 1 - r, s=900, facecolors="none", edgecolors="#B3212D", linewidths=3, zorder=4)
    ax.set_xlim(-0.6, n - 0.4)
    ax.set_ylim(-0.6, n - 0.4)
    ax.set_aspect("equal")
    ax.set_xticks(range(n), [str(i) for i in range(n)])
    ax.set_yticks(range(n), [str(n - 1 - i) for i in range(n)])
    ax.tick_params(length=0, labelsize=9)
    ax.set_title(title, fontsize=11, fontweight="bold", pad=10)
    for spine in ax.spines.values():
        spine.set_visible(False)


def create_figures(results):
    OUT.mkdir(parents=True, exist_ok=True)
    ck = results["connect_k"]
    go = results["go"]

    search_path = OUT / "figure_connect_search.png"
    labels = ["기준 순서", "Renderer 순서", "Renderer 안전축약"]
    values = [
        ck["root_alpha_beta"]["baseline"]["nodes"],
        ck["root_alpha_beta"]["renderer_order"]["nodes"],
        ck["root_alpha_beta"]["renderer_gate"]["nodes"],
    ]
    fig, ax = plt.subplots(figsize=(8.0, 4.2), facecolor="white")
    bars = ax.bar(labels, values, color=["#7B8794", "#3C78A8", "#2E7D5B"], width=0.62)
    ax.set_yscale("log")
    ax.set_ylabel("탐색 노드 수 로그 눈금")
    ax.set_title("같은 최적값을 찾는 데 필요한 탐색량")
    ax.grid(axis="y", color="#D9D9D9", linewidth=0.7, alpha=0.8)
    ax.set_axisbelow(True)
    for bar, value in zip(bars, values):
        ax.text(bar.get_x() + bar.get_width() / 2, value * 1.25, f"{value:,}", ha="center", va="bottom", fontsize=10)
    fig.tight_layout()
    fig.savefig(search_path, dpi=240, bbox_inches="tight", facecolor="white")
    plt.close(fig)

    funnel_path = OUT / "figure_go_residue_counts.png"
    labels = ["보이는 상태", "복수 residue", "합법수 차이"]
    values = [
        go["unique_visible_state_classes"],
        go["classes_with_multiple_reachable_residues"],
        go["classes_with_residue_dependent_legal_actions"],
    ]
    fig, ax = plt.subplots(figsize=(8.0, 4.1), facecolor="white")
    bars = ax.barh(labels[::-1], values[::-1], color=["#B3212D", "#3C78A8", "#7B8794"])
    ax.set_xscale("log")
    ax.set_xlabel("상태 동치류 수 로그 눈금")
    ax.set_title("현재 판이 같아도 residue가 달라지는 상태")
    ax.grid(axis="x", color="#D9D9D9", linewidth=0.7, alpha=0.8)
    ax.set_axisbelow(True)
    for bar, value in zip(bars, values[::-1]):
        ax.text(value * 1.10, bar.get_y() + bar.get_height() / 2, f"{value:,}", va="center", fontsize=10)
    fig.tight_layout()
    fig.savefig(funnel_path, dpi=240, bbox_inches="tight", facecolor="white")
    plt.close(fig)

    witness = go["witness"]
    witness_path = OUT / "figure_go_witness.png"
    fig, axes = plt.subplots(1, 3, figsize=(10.5, 3.5), facecolor="white")
    draw_go_board(axes[0], witness["board_rows"], "같은 현재 판", witness["differing_action"])
    draw_go_board(axes[1], witness["legal_previous_board_rows"], "경로 A 직전 판  합법")
    draw_go_board(axes[2], witness["blocked_previous_board_rows"], "경로 B 직전 판  패로 금지")
    fig.suptitle("백의 0 1 수는 현재 판만으로 판정할 수 없다", fontsize=13, fontweight="bold", y=1.02)
    fig.tight_layout()
    fig.savefig(witness_path, dpi=240, bbox_inches="tight", facecolor="white")
    plt.close(fig)

    state_path = OUT / "figure_state_space.png"
    raw_ck = 2 * 3 ** 16
    raw_go = 3 ** 9 * (3 ** 9 + 1) * 2 * 3
    labels = ["Connect K 원시상한", "Connect K 도달상태", "바둑 원시상한", "바둑 도달상태"]
    vals = [raw_ck, ck["reachable_full_states"], raw_go, go["reachable_full_states"]]
    colors = ["#AAB2BA", "#2E7D5B", "#AAB2BA", "#3C78A8"]
    fig, ax = plt.subplots(figsize=(8.2, 4.3), facecolor="white")
    bars = ax.bar(labels, vals, color=colors, width=0.65)
    ax.set_yscale("log")
    ax.set_ylabel("상태 수 로그 눈금")
    ax.set_title("규칙과 도달가능성이 제거한 원시 상태")
    ax.tick_params(axis="x", rotation=12)
    ax.grid(axis="y", color="#D9D9D9", linewidth=0.7, alpha=0.8)
    ax.set_axisbelow(True)
    for bar, value in zip(bars, vals):
        ax.text(bar.get_x() + bar.get_width()/2, value * 1.18, f"{value:,}", ha="center", va="bottom", fontsize=8.8)
    fig.tight_layout()
    fig.savefig(state_path, dpi=240, bbox_inches="tight", facecolor="white")
    plt.close(fig)

    return {
        "search": search_path,
        "go_counts": funnel_path,
        "witness": witness_path,
        "state_space": state_path,
    }


def build():
    results = json.loads(RESULTS_PATH.read_text(encoding="utf-8"))
    ck = results["connect_k"]
    go = results["go"]
    witness = go["witness"]
    figures = create_figures(results)
    code_hash = hashlib.sha256(CODE_PATH.read_bytes()).hexdigest()
    result_hash = hashlib.sha256(RESULTS_PATH.read_bytes()).hexdigest()

    eq_contract = render_equation(
        r"S_{t+1}=U_D(S_t,\rho_t,a_t;L_D,B_D),\qquad \Phi_t=\mathcal{R}_D(S_t,\rho_t,B_D)",
        "eq_contract.png", width=11.0, fontsize=18,
    )
    eq_sufficiency = render_equation(
        r"M(h_1)=M(h_2)\Rightarrow A(h_1)=A(h_2)\ \wedge\ \forall a\in A,\ M(U(h_1,a))=M(U(h_2,a))",
        "eq_sufficiency.png", width=12.2, fontsize=17,
    )
    eq_lossless = render_equation(
        r"\forall S\in\mathcal{S}_{reach},\quad V_{base}(S)=V_{WRRA}(S),\qquad \sum_S\mathbf{1}[V_{base}(S)\ne V_{WRRA}(S)]=0",
        "eq_lossless.png", width=12.0, fontsize=17,
    )
    eq_go = render_equation(
        r"V_t=(B_t,p_t,c_t),\quad \rho_t=B_{t-1},\quad \exists h_1,h_2:\ V(h_1)=V(h_2)\ \wedge\ A(h_1)\ne A(h_2)",
        "eq_go.png", width=12.0, fontsize=17,
    )
    eq_reduction = render_equation(
        r"\eta=1-\frac{N_{WRRA}}{N_{base}}=1-\frac{32}{245560}=0.9998697",
        "eq_reduction.png", width=8.5, fontsize=18,
    )

    doc = Document()
    configure_document(doc)

    # Cover
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(68)
    set_font(p.add_run("WRRA Game 0 2"), 25, True)
    p = paragraph(doc, "오목형 연결게임 완전탐색과 바둑 residue 상태충분성 검증", align=WD_ALIGN_PARAGRAPH.CENTER, size=15)
    p.paragraph_format.space_after = Pt(34)
    paragraph(doc, "Wonsik Choi  최원식", align=WD_ALIGN_PARAGRAPH.CENTER, size=12)
    paragraph(doc, "janefather@gmail.com", align=WD_ALIGN_PARAGRAPH.CENTER, size=11)
    paragraph(doc, "2026년 9월 17일", align=WD_ALIGN_PARAGRAPH.CENTER, size=11)
    p = paragraph(doc, "WRRA Core 1.0 동결 유지   WRRA Game 0.1 원본 보존", align=WD_ALIGN_PARAGRAPH.CENTER, size=10)
    p.paragraph_format.space_before = Pt(36)
    p = paragraph(doc, "연구 노트와 재현 가능한 계산 장부", align=WD_ALIGN_PARAGRAPH.CENTER, size=10)
    p.paragraph_format.space_before = Pt(20)
    doc.add_page_break()

    heading(doc, "초록", 1)
    paragraph(doc,
        "WRRA Game 0.2는 0.1에서 제시한 오목과 바둑의 구성적 사례를 유한 상태공간의 완전탐색으로 확장한다. "
        "WRRA Core 1.0의 실행문법과 불변축은 수정하지 않았다. 첫 번째 실험은 4×4 판에서 3목을 만드는 Connect K 실험계의 "
        "도달 가능한 6,036,001개 상태를 전부 열거했다. 기준 minimax와 WRRA Renderer의 안전축약 minimax를 모든 상태에서 비교한 결과, "
        "불일치는 0개였다. 빈 판의 alpha beta 탐색은 기준 245,560노드에서 Renderer 순서화 15,686노드, 안전축약 32노드로 줄었고 "
        "최적값은 모두 승리 1로 같았다.")
    paragraph(doc,
        "두 번째 실험은 포획, 자살 금지, 단순 패, 패스, 두 번의 연속 패스 종료를 포함한 3×3 바둑의 전체 도달 그래프를 만들었다. "
        "완전 상태는 132,161개, 합법 전이는 420,710개였다. 현재 판, 차례, 패스 횟수가 같은 72,987개 동치류 가운데 20,538개는 "
        "서로 다른 직전 판 residue를 가졌고, 752개 동치류에서는 그 residue만으로 합법수 집합이 달라졌다. 따라서 단순 패 규칙 아래에서 "
        "현재 판과 차례만으로는 상태가 충분하지 않다. 이 결과는 Renderer가 계산량을 줄이는 모듈이고 residue가 합법수 판정과 상태전이 규칙을 보존하는 "
        "모듈이라는 서로 다른 역할을 정량적으로 분리한다.")

    heading(doc, "핵심 결과", 2)
    table(doc,
        ["시험", "전수 범위", "핵심 수치", "판정"],
        [
            ["Connect K", "도달상태 6,036,001개", "minimax 불일치 0개", "Renderer 축약은 이 실험계에서 무손실"],
            ["Connect K", "빈 판 exact search", "245,560 → 15,686 → 32노드", "순서화와 안전축약 모두 최적값 보존"],
            ["3×3 바둑", "도달상태 132,161개", "residue 의존 동치류 752개", "현재 판만으로 합법수 판정 불충분"],
            ["3×3 바둑", "비종료 상태 107,832개", "패 차단 상태 768개", "각 상태에서 정확히 1수가 residue로 차단"],
        ],
        [1.15, 1.65, 1.65, 2.25],
        center_cols=(0, 1, 2),
    )
    paragraph(doc,
        "이 보고서의 완전성은 명시한 작은 규칙계에 한정된다. 표준 15×15 오목이나 19×19 바둑을 해결했다는 뜻은 아니다. "
        "그러나 WRRA가 단순한 용어 대응을 넘어, 상태충분성과 탐색보존성을 반증 가능한 계산문제로 바꿀 수 있다는 점은 전수검사로 확인되었다.")

    heading(doc, "버전 계보와 연구 질문", 1)
    heading(doc, "0 1에서 0 2로 확장한 부분", 2)
    paragraph(doc,
        "WRRA Game 0.1은 9×9 자유형 오목의 한 구성에서 중앙 한 수가 다음 차례의 승리 출구 4개를 만들고, 가능한 70개 방어 응답 가운데 "
        "그 출구를 모두 없애는 응답이 0개임을 확인했다. 5×5 바둑에서는 같은 현재 판에서 직전 판 residue를 보존하면 합법행동이 18개, "
        "제거하면 19개가 되어 되따내기가 허용되는 사례를 만들었다. 0.2는 이 두 관찰을 전체 상태그래프의 명제로 바꾼다.")
    table(doc,
        ["항목", "WRRA Game 0 1", "WRRA Game 0 2"],
        [
            ["오목 계열", "한 강제승리 구성의 70개 응답 검사", "4×4 Connect K 도달상태 6,036,001개 전수검사"],
            ["바둑", "한 패 사례의 residue 제거시험", "3×3 전체 도달그래프에서 동일 현재 판 동치류 전수비교"],
            ["Renderer", "후보 축약의 구성적 사례", "모든 상태의 minimax 값 보존과 root 노드 감소 측정"],
            ["Residue", "한 되따내기 반례", "752개 동치류와 768개 패 차단 상태 계수"],
        ],
        [1.15, 2.75, 2.8],
        center_cols=(0,),
    )

    heading(doc, "고정한 Core 계약", 2)
    paragraph(doc,
        "상위 규약은 WRRA Core 1.0이다. 이 연구는 게임 규칙을 Domain Profile로 추가했을 뿐, Core의 최소계산, 공통운반자, 표현형, "
        "Fixed Present, Ownership, Boundary, Renderer, Record와 Ledger, Fail Closed 원칙을 바꾸지 않았다.")
    equation(doc, eq_contract, 6.1)
    table(doc,
        ["Core 단계", "Connect K에서의 구현", "바둑에서의 구현"],
        [
            ["SOURCE", "돌 놓기 입력", "착수 또는 패스 입력"],
            ["RELATION LAW", "24개 승리선과 교대 규칙", "인접, 활로, 포획, 자살 금지, 단순 패"],
            ["STATE RESIDUE", "판과 차례  추가 history 없음", "판, 차례, 직전 판, 연속 패스 수"],
            ["BOUNDARY", "살아 있는 승리선과 위협 접점", "돌무리의 활로와 소유 경계"],
            ["COMMON CARRIER", "유한 격자 위 점유 비트", "같은 격자 위 점유 비트와 전이"],
            ["UPDATE", "한 칸 점유 후 차례 교대", "착수, 포획, 패 판정 후 차례 교대"],
            ["RENDERER", "즉시승리, 강제방어, 이중위협, 열린선", "합법수, 패 차단, 포획 가능성"],
            ["PHENOTYPE", "승리, 강제패배, 최적값", "합법 또는 금지, 현재 판의 동일성"],
            ["LEDGER", "상태, 전이, minimax, 노드 수", "상태, 전이, residue별 합법수 집합"],
        ],
        [1.25, 2.65, 2.8],
        center_cols=(0,),
    )

    heading(doc, "상태충분성의 판정 기준", 1)
    paragraph(doc,
        "한 규칙계에서 상태표현 M이 충분하려면, 서로 다른 두 이력 h1과 h2가 같은 M으로 압축될 때 합법행동 집합이 같아야 하고, "
        "같은 행동 뒤의 다음 압축상태도 같아야 한다. 이 조건이 깨지면 서로 다른 합법 전이 구조를 같은 현재로 잘못 합친 것이므로 상태 aliasing이 발생한다.")
    equation(doc, eq_sufficiency, 6.35)
    paragraph(doc,
        "이 정의는 과거 전체를 저장하라는 요구가 아니다. 반대로 다음 전이를 정확히 결정하는 데 필요한 최소 residue만 현재 상태에 남기라는 조건이다. "
        "Connect K에서는 판과 차례만으로 충분하다. 단순 패 바둑에서는 직전 판 하나가 추가로 필요하다. 같은 Core가 두 경우를 구분하는 이유는 "
        "기억을 미리 고정하지 않고 Domain Law가 요구하는 만큼만 보존하기 때문이다.")

    heading(doc, "완전탐색의 의미", 2)
    paragraph(doc,
        "각 실험은 초기 빈 판에서 시작해 합법 전이로 도달할 수 있는 상태만 너비우선으로 닫았다. 원시 조합공간의 모든 배열을 상태로 인정하지 않았다. "
        "턴 순서, 조기 승리, 포획, 자살 금지, 패와 종료 규칙을 통과한 상태만 Ledger에 들어간다. 이 차이는 WRRA의 Boundary가 계산 전에 "
        "존재 가능한 상태와 불가능한 배열을 나누는 과정에 해당한다.")
    figure(doc, figures["state_space"],
        "그림 1  원시 부호배열 상한과 실제 도달상태 수  원시상한은 규칙을 적용하기 전의 조합 수이다", 6.25)
    raw_ck = 2 * 3 ** 16
    raw_go = 3 ** 9 * (3 ** 9 + 1) * 2 * 3
    paragraph(doc,
        f"Connect K의 판과 차례 원시상한은 {raw_ck:,}개이고 실제 도달상태는 {ck['reachable_full_states']:,}개로 {100*ck['reachable_full_states']/raw_ck:.2f}%이다. "
        f"바둑의 현재 판, 직전 판 또는 없음, 차례, 패스 수를 독립 조합한 상한은 {raw_go:,}개지만 실제 도달상태는 {go['reachable_full_states']:,}개로 "
        f"{100*go['reachable_full_states']/raw_go:.5f}%이다. 이 비율은 성능평가가 아니라 규칙이 상태공간을 얼마나 강하게 제한하는지를 보여준다.")

    doc.add_page_break()
    heading(doc, "Connect K 완전탐색", 1)
    heading(doc, "실험 규칙과 범위", 2)
    paragraph(doc,
        "4×4 격자에서 흑이 먼저 두고, 두 사람이 번갈아 빈 칸 하나를 점유한다. 가로, 세로, 두 대각 방향에서 연속 3개 이상을 만들면 즉시 끝난다. "
        "이 규칙은 표준 오목을 축소한 완전열거 실험계다. 목적은 표준 오목의 전략을 해결하는 것이 아니라 Renderer 축약의 정확성을 전체 상태에서 검사하는 것이다.")
    table(doc,
        ["장부 항목", "값"],
        [
            ["승리선 마스크", f"{ck['win_masks']:,}개"],
            ["도달 가능한 완전상태", f"{ck['reachable_full_states']:,}개"],
            ["합법 방향전이", f"{ck['directed_legal_edges']:,}개"],
            ["승리 종료상태", f"{ck['terminal_win_states']:,}개"],
            ["무승부 종료상태", f"{ck['terminal_draw_states']:,}개"],
            ["현재 플레이어 승리값 상태", f"{ck['exact_value_counts_current_player_perspective']['win_1']:,}개"],
            ["무승부값 상태", f"{ck['exact_value_counts_current_player_perspective']['draw_0']:,}개"],
            ["패배값 상태", f"{ck['exact_value_counts_current_player_perspective']['loss_-1']:,}개"],
        ],
        [3.4, 3.25],
        center_cols=(1,),
    )

    heading(doc, "Renderer의 안전축약 규칙", 2)
    paragraph(doc,
        "Renderer는 임의의 가중치로 수를 잘라내지 않았다. 현재 플레이어의 즉시승리 수가 있으면 그 수들만 남긴다. 즉시승리가 없고 상대의 "
        "즉시승리 지점이 하나면 그 지점을 막는 수만 남긴다. 상대의 서로 다른 즉시승리 지점이 둘 이상이면 한 번에 하나의 칸만 점유할 수 있으므로 "
        "현재 상태를 강제패배로 판정한다. 이 세 조건에 해당하지 않을 때는 모든 합법수를 유지하고, 방어, 다음 차례 승리출구 수, 열린 승리선의 "
        "점유도와 중심성을 이용해 탐색 순서만 정한다.")
    table(doc,
        ["Renderer 표현형", "상태 수", "계산상 의미"],
        [
            ["즉시승리", f"{ck['phenotype_census']['immediate_win']:,}", "승리수만 평가해도 값 1이 확정"],
            ["단일 강제방어", f"{ck['phenotype_census']['mandatory_block']:,}", "유일한 위협 칸만 다음 상태 후보"],
            ["다중 위협 강제패배", f"{ck['phenotype_census']['forced_loss_certificate']:,}", "추가 전개 없이 값 -1을 인증"],
            ["열린 탐색", f"{ck['phenotype_census']['open_search']:,}", "모든 합법수를 보존하고 순서만 변경"],
            ["종료상태", f"{ck['phenotype_census']['terminal_win'] + ck['phenotype_census']['terminal_draw']:,}", "이미 승패 또는 무승부 확정"],
        ],
        [2.0, 1.35, 3.3],
        center_cols=(0, 1),
    )

    heading(doc, "모든 상태에서 최적값 보존", 2)
    paragraph(doc,
        "기준 solver와 안전축약 solver를 도달 가능한 모든 상태에서 각각 실행했다. 두 solver는 현재 플레이어 관점의 승리 1, 무승부 0, 패배 -1을 반환한다. "
        "6,036,001개 상태에서 값이 다른 경우는 한 건도 없었다. 따라서 이 유한 규칙계에서는 Renderer가 단순한 휴리스틱이 아니라 최적값을 보존하는 "
        "계산 축약이라는 명제가 성립한다.")
    equation(doc, eq_lossless, 6.35)
    table(doc,
        ["비교", "검사 상태", "불일치", "판정"],
        [["기준 minimax 대 안전축약 minimax", f"{ck['reachable_full_states']:,}", f"{ck['all_state_value_mismatches']:,}", "PASS"]],
        [3.2, 1.35, 1.0, 1.1],
        center_cols=(1, 2, 3),
    )

    heading(doc, "탐색량 감소", 2)
    paragraph(doc,
        "같은 빈 판에서 transposition cache 없이 alpha beta를 실행해 수순 정렬과 안전축약의 효과를 분리했다. 기준 행 우선 순서는 245,560노드를 방문했다. "
        "모든 수를 유지한 채 Renderer가 순서만 바꾸면 15,686노드로 93.61% 줄었다. 안전조건으로 후보까지 줄이면 32노드가 되었고 기준 대비 "
        "99.98697% 감소, 7,673.75배의 차이를 보였다. 세 실행의 root 최적값은 모두 1이다.")
    figure(doc, figures["search"], "그림 2  빈 판 alpha beta의 방문 노드 수  세 탐색의 최적값은 모두 1이다", 6.2)
    equation(doc, eq_reduction, 4.7)
    paragraph(doc,
        "32노드라는 수는 보편적인 오목 성능을 뜻하지 않는다. 4×4 3목의 위협 구조가 매우 빨리 갈라지기 때문에 가능한 결과다. 학문적으로 중요한 수치는 "
        "감소율 자체보다 6,036,001개 전 상태에서 값 보존을 먼저 확인한 뒤 그 축약을 사용했다는 순서다.")

    doc.add_page_break()
    heading(doc, "바둑 residue 완전탐색", 1)
    heading(doc, "규칙과 완전상태", 2)
    paragraph(doc,
        "3×3 격자에서 직교 인접을 사용하고, 활로가 없는 상대 돌무리를 포획하며, 자살수는 금지한다. 새 판이 직전 판과 같아지는 즉시 되따내기는 단순 패로 금지한다. "
        "패스는 항상 가능하고 두 번의 연속 패스에서 종료한다. 완전상태는 현재 판 B, 둘 차례 p, 직전 판 residue, 연속 패스 수 c로 정의했다.")
    table(doc,
        ["장부 항목", "값"],
        [
            ["도달 가능한 완전상태", f"{go['reachable_full_states']:,}개"],
            ["합법 방향전이", f"{go['directed_legal_edges']:,}개"],
            ["착수 전이", f"{go['placement_edges']:,}개"],
            ["패스 전이", f"{go['pass_edges']:,}개"],
            ["두 패스 종료상태", f"{go['terminal_two_pass_states']:,}개"],
            ["현재 판 차례 패스 수 동치류", f"{go['unique_visible_state_classes']:,}개"],
        ],
        [3.5, 3.15],
        center_cols=(1,),
    )

    heading(doc, "현재 판만으로 충분한가", 2)
    paragraph(doc,
        "직전 판만의 영향을 분리하기 위해 현재 판, 차례, 연속 패스 수가 같은 상태들을 하나의 동치류로 묶었다. 그 안에서 직전 판이 여러 개 존재하는지, "
        "그리고 직전 판에 따라 합법 착수 집합이 달라지는지를 비교했다. 패스는 모든 비종료 상태에서 동일하게 가능하므로 합법수 차이는 착수에 대해서만 셌다.")
    equation(doc, eq_go, 6.35)
    figure(doc, figures["go_counts"], "그림 3  보이는 현재가 같을 때 residue 다양성과 합법수 차이", 6.15)
    table(doc,
        ["분류", "개수", "해석"],
        [
            ["보이는 상태 동치류", f"{go['unique_visible_state_classes']:,}", "현재 판, 차례, 패스 수가 같은 집합"],
            ["복수 residue 동치류", f"{go['classes_with_multiple_reachable_residues']:,}", f"전체의 {100*go['classes_with_multiple_reachable_residues']/go['unique_visible_state_classes']:.2f}%"],
            ["합법수 차이 동치류", f"{go['classes_with_residue_dependent_legal_actions']:,}", f"전체의 {100*go['classes_with_residue_dependent_legal_actions']/go['unique_visible_state_classes']:.2f}%"],
            ["합법수 차이 동치류의 완전상태", f"{go['full_states_in_action_differing_classes']:,}", "같은 현재 판이라도 합법수 집합이 하나로 정해지지 않음"],
            ["실제 패 차단 완전상태", f"{go['ko_sensitive_full_states']:,}", f"비종료 상태의 {100*go['ko_sensitive_full_states']/(go['reachable_full_states']-go['terminal_two_pass_states']):.2f}%"],
            ["한 현재 판의 최대 residue", f"{go['max_reachable_residues_per_visible_state']:,}", "서로 다른 직전 판의 최대 수"],
            ["한 현재 판의 최대 합법수 집합", f"{go['max_distinct_legal_action_sets_per_visible_state']:,}", "같은 판에 최대 세 가지 합법수 집합이 대응"],
        ],
        [2.25, 1.25, 3.15],
        center_cols=(0, 1),
    )
    paragraph(doc,
        "패 차단 상태 768개는 전체에서 드물다. 그러나 상태충분성은 빈도 문제가 아니다. 단 하나의 도달 가능한 반례만 있어도 현재 판만으로 다음 합법행동을 "
        "결정할 수 없다는 명제는 무너진다. 여기서는 반례가 752개 동치류에 걸쳐 반복되었고, 각 패 차단 상태에서 정확히 한 수가 residue 때문에 제외되었다.")

    heading(doc, "가장 짧은 두 경로의 반례", 2)
    paragraph(doc,
        "완전탐색이 찾은 가장 짧은 증명쌍은 두 경로 모두 5수다. 두 경로는 같은 현재 판 B.B / WB. / ...과 백 차례, 패스 수 0에 도달한다. "
        "경로 A에서는 백의 0 1 착수가 합법이고, 경로 B에서는 같은 수가 직전 판을 복원하므로 패로 금지된다. 좌표는 위에서 아래로 행 0부터, "
        "왼쪽에서 오른쪽으로 열 0부터 센다.")
    figure(doc, figures["witness"], "그림 4  같은 현재 판과 서로 다른 직전 판  붉은 원은 백의 시험수 0 1이다", 6.45)
    legal_history = "  ".join(
        f"{m['ply']} {m['player']} {m['action'] if isinstance(m['action'], str) else tuple(m['action'])}"
        for m in witness["legal_history"]
    )
    blocked_history = "  ".join(
        f"{m['ply']} {m['player']} {m['action'] if isinstance(m['action'], str) else tuple(m['action'])}"
        for m in witness["blocked_history"]
    )
    table(doc,
        ["경로", "5수 이력", "백의 0 1"],
        [
            ["A", legal_history, "합법"],
            ["B", blocked_history, "패로 금지"],
        ],
        [0.65, 5.05, 1.0],
        center_cols=(0, 2),
    )
    paragraph(doc,
        "경로 B에서 백이 0 1에 두면 B의 0 0 돌이 잡히고, 새 판은 직전 판 .WB / WB. / ...과 정확히 같아진다. 경로 A의 직전 판은 "
        "B.B / W.. / ...이므로 같은 착수가 과거 판을 복원하지 않는다. 이 한 쌍은 판의 모양과 차례가 같다는 사실만으로 합법성을 판정할 수 없음을 직접 보여준다.")

    doc.add_page_break()
    heading(doc, "두 실험을 함께 볼 때 드러나는 구조", 1)
    heading(doc, "Renderer와 residue의 역할은 다르다", 2)
    table(doc,
        ["제거시험", "관측된 변화", "모듈의 역할"],
        [
            ["Connect K Renderer 순서 제거", "15,686 → 245,560노드", "정확도를 바꾸지 않고 탐색비용을 줄임"],
            ["Connect K 안전축약 제거", "32 → 245,560노드", "국소 표현형을 계산 인증서로 사용"],
            ["바둑 직전 판 residue 제거", "768개 상태에서 금지수 1개가 합법수로 유입", "합법수 판정과 다음 상태의 전이 규칙을 보존"],
            ["전체 이력 대신 최소 residue 사용", "단순 패에는 직전 판 하나로 충분", "필요한 기억만 현재 상태에 결합"],
        ],
        [2.15, 2.0, 2.55],
        center_cols=(0,),
    )
    paragraph(doc,
        "Renderer는 상태를 바꾸지 않고 결정에 필요한 구조를 드러낸다. residue는 보이는 상태에 없는 전이 제약을 현재로 운반한다. 하나는 계산량을 줄이고, "
        "다른 하나는 합법성 자체를 보존한다. 두 모듈을 한 단어로 묶지 않고 제거시험으로 역할을 분리할 수 있다는 점이 이번 확장의 중요한 결과다.")

    heading(doc, "Fixed Present와 최소기억", 2)
    paragraph(doc,
        "Fixed Present는 과거가 존재하지 않는다는 뜻이 아니라, 다음 계산에 필요한 과거의 영향이 현재 상태의 residue로 완결되어야 한다는 뜻으로 구현된다. "
        "Connect K는 추가 residue가 0이다. 단순 패 바둑은 직전 판 1개가 필요하다. 위치 전체 반복을 금지하는 positional superko를 사용한다면 방문 판 집합이 "
        "필요하므로 residue가 커진다. 필요한 기억량은 Core의 취향이 아니라 Domain Law에서 유도된다.")

    heading(doc, "공통운반자의 실제 의미", 2)
    paragraph(doc,
        "두 게임의 세부 법칙은 다르지만, 격자 위 점유상태, 합법행동, 결정적 update, 경계, residue, Renderer, phenotype, Ledger라는 같은 실행 인터페이스로 "
        "표현된다. 공통운반자는 서로 다른 현상을 같은 현상이라고 부르는 장치가 아니다. 각 법칙의 차이를 보존하면서도 동일한 계산 질문을 던질 수 있게 하는 "
        "최소 형식이다. 이번 연구에서는 그 공통 질문이 도달 가능한 상태인가, 다음 행동이 합법인가, 현재 표현만으로 다음 합법 전이 집합을 결정할 수 있는가, 표현형이 최적값을 "
        "보존하면서 후보를 줄이는가였다.")

    heading(doc, "새로 확인된 학문적 장점", 1)
    numbered(doc, [
        "상태충분성을 설계자의 직관이 아니라 동치류 반례로 판정할 수 있다. 같은 현재 표현에 서로 다른 합법수 집합이 나타나면 residue 부족이 자동으로 드러난다.",
        "표현형을 설명용 라벨에서 계산 인증서로 바꿀 수 있다. Connect K의 즉시승리, 단일 강제방어, 다중 위협은 값이 보존되는 후보축약 규칙이 된다.",
        "정확성과 효율성을 서로 다른 제거시험으로 분리할 수 있다. residue 제거는 잘못된 수를 허용했고, Renderer 제거는 답을 바꾸지 않은 채 탐색량을 늘렸다.",
        "Markovian 규칙과 history dependent 규칙을 같은 Core 안에서 다룰 수 있다. 차이는 Core를 고치는 대신 Domain Profile의 최소 residue 크기로 표현된다.",
        "원시 상태공간과 도달 상태공간을 구분해 계산 낭비의 원인을 측정할 수 있다. 법과 경계를 먼저 적용하면 존재 불가능한 배열을 탐색에서 제외한다.",
        "모든 수치를 Ledger로 남겨 다른 연구자가 동일한 상태 수, 전이 수, 반례 이력, minimax 값을 다시 계산할 수 있다.",
    ])
    paragraph(doc,
        "이 장점들은 게임을 잘 둔다는 주장과 별개다. 더 일반적인 가치는 복잡계의 표현이 충분한지, 어떤 기억이 필요한지, 어떤 phenotype이 손실 없이 계산을 "
        "줄이는지를 하나의 검증 절차로 바꾼 데 있다. 전력망, 생물학적 조절, 시장 규칙처럼 현재 관측이 같아도 숨은 residue에 따라 합법 전이 판정이 달라지는 "
        "영역에서도 같은 상태 aliasing 시험을 사용할 수 있다.")

    heading(doc, "판정 범위와 다음 단계", 1)
    heading(doc, "이번에 확정된 것", 2)
    bullets(doc, [
        "4×4 3목의 전체 도달상태 6,036,001개에서 안전축약 minimax와 기준 minimax의 값이 같았다.",
        "Renderer의 순서화만으로 빈 판 alpha beta 노드가 93.61% 감소했고, 증명된 안전축약을 더하면 99.98697% 감소했다.",
        "3×3 단순 패 바둑의 전체 도달상태 132,161개 가운데 768개에서 직전 판 residue가 정확히 한 착수를 차단했다.",
        "같은 현재 판, 차례, 패스 수를 가진 두 실제 이력이 서로 다른 합법수 집합을 가질 수 있으므로 board only 표현은 불충분하다.",
        "WRRA Core 1.0은 수정되지 않았고, WRRA Game 0.1도 원본으로 보존된다.",
    ])

    heading(doc, "아직 확정하지 않은 것", 2)
    bullets(doc, [
        "표준 15×15 오목과 19×19 바둑의 완전해 또는 실전 기력",
        "전문화된 게임엔진보다 우수하다는 성능 주장",
        "Renderer 우선순위가 모든 판 크기와 규칙에서 같은 감소율을 보인다는 일반화",
        "단순 패가 아닌 positional superko에서 필요한 최소 residue의 압축형",
        "학습 기반 평가함수와 결합했을 때의 정확도 및 계산비용",
    ])
    paragraph(doc,
        "다음 버전에서는 5×5 또는 6×6 연결게임의 대칭축약, 위협깊이별 증명서, 바둑의 positional superko와 visited set 압축, 그리고 장기나 체스처럼 "
        "이동과 포획이 함께 있는 게임에서 상태충분성 검사를 추가할 수 있다. 다음 단계의 기준도 동일하다. 새 용어를 더하는 것보다 제거시험과 전수 또는 경계가 "
        "명확한 계산으로 Core 구성요소의 필요성을 확인한다.")

    heading(doc, "재현 장부", 1)
    paragraph(doc,
        "계산은 정수 비트보드와 결정적인 행 우선 수순을 사용한다. Connect K는 초기 빈 판에서 조기 승리 시 확장을 멈추고 전체 도달그래프를 닫았다. "
        "바둑은 완전상태를 현재 흑돌 비트, 백돌 비트, 차례, 직전 흑돌 비트, 직전 백돌 비트, 연속 패스 수로 저장했다. 같은 완전상태를 다시 만나면 "
        "확장하지 않았으며, 큐가 빌 때까지 진행했다.")
    table(doc,
        ["파일", "역할", "SHA 256"],
        [
            [CODE_PATH.name, "전수열거와 검증 코드", code_hash],
            [RESULTS_PATH.name, "기계판독 결과 장부", result_hash],
        ],
        [2.3, 1.8, 2.6],
        center_cols=(0,),
    )
    paragraph(doc,
        f"실행시간 장부는 Connect K {ck['elapsed_seconds']:.2f}초, 바둑 {go['elapsed_seconds']:.2f}초, 전체 {results['elapsed_seconds_total']:.2f}초다. "
        "실행시간은 환경에 따라 변하지만 상태 수, 전이 수, minimax 값, 동치류 수와 반례 이력은 결정적이다.")

    heading(doc, "재현 절차", 2)
    numbered(doc, [
        "Python 3 환경에서 wrra_game_0_2.py를 실행하고 JSON 출력 경로를 지정한다.",
        "Connect K 결과에서 reachable full states 6,036,001과 all state value mismatches 0을 확인한다.",
        "바둑 결과에서 reachable full states 132,161과 residue dependent legal actions classes 752를 확인한다.",
        "witness의 두 5수 이력을 재생하고, 같은 현재 판에서 백의 0 1이 한 경로에서는 합법이고 다른 경로에서는 직전 판 복원으로 금지되는지 확인한다.",
    ])

    heading(doc, "참조 계보", 2)
    paragraph(doc, "WRRA Core 1.0  Frozen Domain Agnostic Execution Architecture  DOI 10.5281/zenodo.22650956")
    paragraph(doc, "WRRA Game 0.1  오목과 바둑 복잡계 분석  2026년 9월 16일")
    paragraph(doc, "WRRA Game 0.2  본 보고서와 재현 코드  2026년 9월 17일")

    heading(doc, "최종 판정", 1)
    paragraph(doc,
        "WRRA Game 0.2에서 가장 강한 결과는 두 게임을 같은 말로 설명했다는 사실이 아니다. 같은 실행계약으로 서로 반대되는 기억 요구를 정확히 분리했다는 점이다. "
        "Connect K는 현재 판과 차례만으로 충분했고, Renderer는 최적값을 잃지 않으면서 계산을 줄였다. 단순 패 바둑은 같은 현재 판에 서로 다른 합법수 집합이 대응할 수 "
        "있었고, 직전 판 residue가 그 전이 차이를 보존했다. 최소계산은 무조건 적게 저장하거나 적게 탐색하는 것이 아니라, 결과를 바꾸지 않는 범위까지만 줄이는 원칙으로 "
        "구현되었다.")
    paragraph(doc,
        "따라서 이번 버전은 WRRA의 게임 적용을 구성적 사례에서 유한 상태공간의 검증으로 한 단계 전진시킨다. 명시한 작은 규칙계 안에서는 Renderer의 무손실 축약과 "
        "residue의 필수성이 모두 수치와 반례 이력으로 재현 가능하다.")

    props = doc.core_properties
    props.title = "WRRA Game 0 2 오목형 연결게임 완전탐색과 바둑 residue 상태충분성 검증"
    props.subject = "WRRA Core 1.0 Domain Profile exhaustive validation"
    props.author = "Wonsik Choi"
    props.keywords = "WRRA, Connect K, Omok, Go, minimax, residue, renderer, state sufficiency"
    props.comments = "WRRA Core 1.0 frozen; WRRA Game 0.1 preserved"

    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build()
