#!/usr/bin/env python3
"""Build the integrated Korean WRRA Game 1.0 research report."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path

import matplotlib.pyplot as plt
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from build_wrra_game_0_2_report import (
    MATH_FONT,
    configure_document,
    create_figures,
    draw_go_board,
    equation,
    figure,
    heading,
    numbered,
    paragraph,
    set_font,
    table,
    bullets,
)


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "wrra_game_output" / "v1_0"
DOCX_PATH = ROOT / "WRRA_Game_1.0_복잡게임의_최소상태와_실제_바둑엔진_최원식_2026-09-17.docx"
R01_PATH = ROOT / "wrra_game_output" / "wrra_game_0_1_results.json"
R02_PATH = ROOT / "wrra_game_output" / "wrra_game_0_2_results.json"
BENCH_PATH = ROOT / "wrra_go_output" / "wrra_go_0_1_benchmark.json"
TEST_PATH = ROOT / "wrra_go_output" / "wrra_go_0_1_rule_tests.json"
SELFPLAY_PATH = ROOT / "wrra_go_output" / "wrra_go_0_1_selfplay_9x9.json"
CODE_PATHS = [ROOT / "wrra_game_0_1.py", ROOT / "wrra_game_0_2.py", ROOT / "wrra_go_0_1.py"]
NAVY = "#24364B"
BLUE = "#3C78A8"
GREEN = "#2E7D5B"
RED = "#B3212D"


def file_hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def render_equation(tex: str, filename: str, width=10.0, height=0.75, fontsize=18) -> Path:
    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / filename
    fig = plt.figure(figsize=(width, height), facecolor="white")
    fig.text(0.5, 0.5, f"${tex}$", ha="center", va="center", fontsize=fontsize, fontfamily=MATH_FONT)
    fig.savefig(path, dpi=260, bbox_inches="tight", pad_inches=0.08, facecolor="white")
    plt.close(fig)
    return path


def create_engine_figures(benchmark: dict, selfplay: dict) -> dict[str, Path]:
    OUT.mkdir(parents=True, exist_ok=True)

    architecture = OUT / "figure_wrra_go_architecture.png"
    fig, ax = plt.subplots(figsize=(10.5, 3.8), facecolor="white")
    ax.set_xlim(0, 10.5)
    ax.set_ylim(0, 3.8)
    ax.axis("off")
    labels = [
        (0.25, "완전상태\n판 차례 residue 패스", "#E9EEF3"),
        (2.35, "합법전이\n포획 자살금지 패", "#E7F1EA"),
        (4.45, "Renderer\n11개 투명 채널", "#E7EEF6"),
        (6.55, "MCTS\n허용 분기 평가", "#F6EFE3"),
        (8.65, "착수와 Ledger\n근거를 함께 기록", "#F4E8EA"),
    ]
    for x, label, colour in labels:
        patch = plt.Rectangle((x, 1.15), 1.6, 1.25, facecolor=colour, edgecolor="#66717C", linewidth=1.2)
        ax.add_patch(patch)
        ax.text(x + 0.8, 1.77, label, ha="center", va="center", fontsize=10.5, fontweight="bold")
    for x in (1.88, 3.98, 6.08, 8.18):
        ax.annotate("", xy=(x + 0.38, 1.77), xytext=(x, 1.77), arrowprops=dict(arrowstyle="->", color="#24364B", lw=1.8))
    ax.text(5.25, 3.12, "WRRA Go 0.1의 한 수 결정 과정", ha="center", fontsize=14, fontweight="bold")
    ax.text(5.25, 0.55, "결과는 미래 예측이 아니라 현재 규칙에서 허용되는 후속 상태의 탐색과 가치평가이다", ha="center", fontsize=10.5)
    fig.tight_layout()
    fig.savefig(architecture, dpi=240, bbox_inches="tight", facecolor="white")
    plt.close(fig)

    bench = OUT / "figure_wrra_go_benchmark.png"
    keys = ["uniform_mcts", "random"]
    labels = ["동일 계산량\n균등 MCTS", "무작위 착수"]
    summaries = [benchmark["summary"][key] for key in keys]
    rates = [summary["wrra_win_rate"] for summary in summaries]
    lower = [rate - summary["wilson_95_interval"][0] for rate, summary in zip(rates, summaries)]
    upper = [summary["wilson_95_interval"][1] - rate for rate, summary in zip(rates, summaries)]
    fig, ax = plt.subplots(figsize=(7.8, 4.6), facecolor="white")
    bars = ax.bar(labels, rates, color=[BLUE, GREEN], width=0.58, yerr=[lower, upper], capsize=7)
    ax.set_ylim(0, 1.08)
    ax.set_ylabel("WRRA MCTS 승률")
    ax.set_title(f"{benchmark['rules']['board']} 색상 교대 대국  각 대조군 {summaries[0]['games']}국")
    ax.grid(axis="y", color="#D9D9D9", linewidth=0.7)
    ax.set_axisbelow(True)
    for bar, rate, summary in zip(bars, rates, summaries):
        ax.text(bar.get_x() + bar.get_width() / 2, min(1.035, rate + 0.035), f"{summary['wrra_wins']} / {summary['games']}", ha="center", fontsize=11, fontweight="bold")
    fig.tight_layout()
    fig.savefig(bench, dpi=240, bbox_inches="tight", facecolor="white")
    plt.close(fig)

    final_board = OUT / "figure_wrra_go_selfplay_final.png"
    fig, ax = plt.subplots(figsize=(5.5, 5.5), facecolor="white")
    draw_go_board(ax, selfplay["final_board"], "9×9 WRRA 자기대국 최종 판")
    fig.tight_layout()
    fig.savefig(final_board, dpi=240, bbox_inches="tight", facecolor="white")
    plt.close(fig)

    candidates = OUT / "figure_wrra_go_first_move.png"
    first = selfplay["moves"][0]["decision"]["candidates"][:8]
    labels = [item["action"] for item in first]
    visits = [item["visits"] for item in first]
    colours = [GREEN if i == 0 else "#AAB2BA" for i in range(len(first))]
    fig, ax = plt.subplots(figsize=(8.0, 4.2), facecolor="white")
    bars = ax.bar(labels, visits, color=colours)
    ax.set_ylabel("MCTS 방문 수")
    ax.set_title("9×9 자기대국 첫 수의 후보 장부")
    ax.grid(axis="y", color="#D9D9D9", linewidth=0.7)
    ax.set_axisbelow(True)
    for bar, value in zip(bars, visits):
        ax.text(bar.get_x() + bar.get_width() / 2, value + 0.2, str(value), ha="center", fontsize=9)
    fig.tight_layout()
    fig.savefig(candidates, dpi=240, bbox_inches="tight", facecolor="white")
    plt.close(fig)

    return {
        "architecture": architecture,
        "benchmark": bench,
        "final_board": final_board,
        "candidates": candidates,
    }


def add_cover(doc: Document):
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(52)
    set_font(p.add_run("WRRA Game 1 0"), 25, True)
    p = paragraph(doc, "복잡게임의 최소 충분상태와 실제 바둑 엔진", align=WD_ALIGN_PARAGRAPH.CENTER, size=16)
    p.paragraph_format.space_after = Pt(12)
    p = paragraph(doc, "오목과 바둑의 구성적 분석 완전탐색 검증 및 WRRA Go 0 1", align=WD_ALIGN_PARAGRAPH.CENTER, size=12)
    p.paragraph_format.space_after = Pt(34)
    paragraph(doc, "Wonsik Choi  최원식", align=WD_ALIGN_PARAGRAPH.CENTER, size=12)
    paragraph(doc, "Independent Researcher", align=WD_ALIGN_PARAGRAPH.CENTER, size=10.5)
    paragraph(doc, "janefather@gmail.com", align=WD_ALIGN_PARAGRAPH.CENTER, size=11)
    paragraph(doc, "2026년 9월 17일", align=WD_ALIGN_PARAGRAPH.CENTER, size=11)
    p = paragraph(doc, "WRRA Core 1.0 동결 유지", align=WD_ALIGN_PARAGRAPH.CENTER, size=10)
    p.paragraph_format.space_before = Pt(30)
    paragraph(doc, "Zenodo 공개용 통합 연구보고서와 재현 패키지", align=WD_ALIGN_PARAGRAPH.CENTER, size=10)
    doc.add_page_break()


def build():
    r01 = json.loads(R01_PATH.read_text(encoding="utf-8"))
    r02 = json.loads(R02_PATH.read_text(encoding="utf-8"))
    benchmark = json.loads(BENCH_PATH.read_text(encoding="utf-8"))
    tests = json.loads(TEST_PATH.read_text(encoding="utf-8"))
    selfplay = json.loads(SELFPLAY_PATH.read_text(encoding="utf-8"))
    ck = r02["connect_k"]
    go = r02["go"]
    figures02 = create_figures(r02)
    engine_figures = create_engine_figures(benchmark, selfplay)

    eq_state = render_equation(
        r"S_t=(B_t,p_t,\rho_t,c_t,n_t),\qquad \rho_t=B_{t-1}",
        "eq_engine_state.png", width=8.5,
    )
    eq_update = render_equation(
        r"T(S_t,a)=S_{t+1}\quad\text{iff}\quad a\in A_{legal}(S_t;L_D,\rho_t)",
        "eq_engine_update.png", width=10.0,
    )
    eq_renderer = render_equation(
        r"R(a\mid S)=\sum_{i=1}^{11}w_i f_i(S,a)",
        "eq_engine_renderer.png", width=7.2,
    )
    eq_puct = render_equation(
        r"a^*=\arg\max_a\left(Q(S,a)+c\,P_R(a\mid S)\frac{\sqrt{N(S)}}{1+N(S,a)}\right)",
        "eq_engine_puct.png", width=11.0,
    )
    eq_suff = render_equation(
        r"M(h_1)=M(h_2)\Rightarrow A_{legal}(h_1)=A_{legal}(h_2)",
        "eq_engine_suff.png", width=9.4,
    )
    eq_lossless = render_equation(
        r"\sum_{S\in\mathcal{S}_{reach}}\mathbf{1}[V_{base}(S)\ne V_{WRRA}(S)]=0",
        "eq_engine_lossless.png", width=9.5,
    )

    doc = Document()
    configure_document(doc)
    add_cover(doc)

    heading(doc, "연구 결론", 1)
    paragraph(doc,
        "WRRA Game 연구는 하나의 고정된 실행문법을 오목형 연결게임과 바둑에 적용해 세 단계를 완성했다. 0.1은 강제승리와 패의 구성적 사례를 계산했고, "
        "0.2는 두 축소 규칙계의 전체 도달상태를 열거해 Renderer의 값 보존과 residue의 필수성을 검증했으며, 1.0은 같은 상태와 전이 구조 위에 실제 착수를 "
        "선택하는 WRRA Go 0.1 엔진을 구현했다. 따라서 현재 결과는 개념 대응, 전수검증, 실행 가능한 의사결정 알고리즘을 하나의 계보로 연결한다.")
    table(doc,
        ["단계", "질문", "실행 결과", "확정된 의미"],
        [
            ["Game 0.1", "구조가 실제 사례를 계산하는가", "9×9 오목 70응답과 5×5 패 반례", "구성적 작동 가능성"],
            ["Game 0.2", "값과 합법성이 전 상태에서 보존되는가", "6,036,001 Connect K 상태와 132,161 바둑 상태", "유한 규칙계의 전수검증"],
            ["Game 1.0", "같은 구조로 실제 착수를 선택할 수 있는가", "9×9 대국 가능 엔진과 색상 교대 벤치마크", "계산 구조에서 행위자로 확장"],
        ],
        [1.05, 2.0, 2.2, 1.65], center_cols=(0,),
    )
    paragraph(doc,
        "가장 중요한 용어도 바로잡았다. residue는 미래를 예측하거나 보존하지 않는다. residue는 현재 판만으로 판정할 수 없는 합법수 제약을 현재 상태에 포함해 "
        "합법수 집합과 다음 상태의 갱신 규칙을 정확히 만든다. 탐색 역시 이미 존재하는 미래를 알아내는 과정이 아니라, 현재 규칙이 허용하는 후속 상태를 생성하고 "
        "목표함수에 따라 비교하는 계산이다.")

    doc.add_page_break()
    heading(doc, "핵심 수치", 2)
    uniform = benchmark["summary"]["uniform_mcts"]
    random_summary = benchmark["summary"]["random"]
    table(doc,
        ["영역", "전수 또는 대국 범위", "결과", "판정"],
        [
            ["Connect K", f"도달상태 {ck['reachable_full_states']:,}개", f"minimax 불일치 {ck['all_state_value_mismatches']:,}개", "안전축약 값 보존"],
            ["Connect K", "빈 판 alpha beta", f"{ck['root_alpha_beta']['baseline']['nodes']:,} → {ck['root_alpha_beta']['renderer_gate']['nodes']:,}노드", "같은 root 값"],
            ["3×3 바둑", f"도달상태 {go['reachable_full_states']:,}개", f"residue 의존 동치류 {go['classes_with_residue_dependent_legal_actions']:,}개", "board only 불충분"],
            ["WRRA Go", f"균등 MCTS 상대 {uniform['games']}국", f"{uniform['wrra_wins']}승 {uniform['wrra_losses']}패", "같은 simulation 수"],
            ["WRRA Go", f"무작위 상대 {random_summary['games']}국", f"{random_summary['wrra_wins']}승 {random_summary['wrra_losses']}패", "색상 교대"],
            ["9×9 자기대국", f"{selfplay['move_count']}수", f"{selfplay['winner']} 승  흑 기준 {selfplay['score']['black_margin']:+.1f}집", selfplay["ended_by"]],
        ],
        [1.25, 1.9, 1.95, 1.75], center_cols=(0, 1, 2),
    )
    paragraph(doc,
        "대국 수치는 전문 바둑엔진과의 기력 비교가 아니라, WRRA Renderer가 실제 합법수 생성과 수 선택에 연결되고 동일 계산량의 통제 알고리즘과 비교 가능한지를 "
        "확인한 첫 실행시험이다. 모든 대국은 색을 교대했고 같은 판 크기, 같은 simulation 수, 같은 득점 규칙을 사용했다.")

    heading(doc, "문서 구성", 2)
    numbered(doc, [
        "WRRA Game 공통 실행계약과 예측으로 오인될 표현의 정정",
        "Game 0.1의 오목 강제승리와 바둑 패 구성적 분석",
        "Game 0.2의 Connect K 및 3×3 바둑 완전탐색",
        "WRRA Go 0.1의 상태, Renderer, MCTS 결합과 실제 코드",
        "규칙검사, 대조 대국, 9×9 자기대국 및 재현 장부",
        "학문적 가치, 적용 범위, 후속 개발 기준과 Zenodo 메타데이터",
    ])

    doc.add_page_break()
    heading(doc, "1 공통 실행계약", 1)
    heading(doc, "WRRA Core 1 0을 고치지 않는 응용", 2)
    paragraph(doc,
        "게임은 WRRA Core 1.0의 수정안이 아니라 Domain Profile이다. Core는 어떤 게임을 특별 취급하지 않고 SOURCE, LAW, STATE와 RESIDUE, BOUNDARY, "
        "COMMON CARRIER, UPDATE, RENDERER, PHENOTYPE, LEDGER의 실행 순서만 고정한다. 각 게임의 규칙과 목표는 Domain Profile에 들어간다. 이 분리는 "
        "오목과 바둑의 차이를 지우지 않으면서 같은 분석 질문을 던지게 한다.")
    table(doc,
        ["Core 항목", "오목형 연결게임", "바둑", "바둑 엔진"],
        [
            ["SOURCE", "빈 칸의 착수", "착수 또는 패스", "사람 또는 알고리즘의 행동"],
            ["LAW", "교대와 연속 K개", "활로 포획 자살금지 단순 패", "같은 규칙의 실행 함수"],
            ["STATE", "판과 차례", "판 차례 residue 패스 수", "수 번호와 탐색노드 추가"],
            ["RESIDUE", "추가 이력 0", "직전 판 1개", "패 판정에 그대로 사용"],
            ["BOUNDARY", "빈 칸과 승리선", "합법 착수 경계", "합법 후보만 탐색"],
            ["CARRIER", "유한 격자", "유한 직교 격자", "동일 좌표와 점유 배열"],
            ["UPDATE", "돌 추가와 차례교대", "포획 패 판정 후 교대", "불변 GoState 반환"],
            ["RENDERER", "승리 위협 강제방어", "포획 활로 연결 절단", "prior와 rollout 순서"],
            ["PHENOTYPE", "승리 강제패배 열린탐색", "capture rescue atari 등", "착수 설명 라벨"],
            ["LEDGER", "상태 값 노드", "합법수와 반례 이력", "후보 방문 가치 대국기록"],
        ],
        [1.05, 1.75, 1.85, 1.95], center_cols=(0,),
    )

    heading(doc, "상태충분성", 2)
    paragraph(doc,
        "상태표현 M이 충분하려면 서로 다른 두 합법 이력이 같은 M으로 압축될 때 합법행동 집합이 같아야 한다. 같은 판 모양을 보고도 이력에 따라 허용되는 수가 "
        "달라진다면, M에는 현재의 다음 계산에 필요한 정보가 빠져 있다. 이것이 residue를 추가하는 기준이며 과거 전체를 무조건 저장하는 기준이 아니다.")
    equation(doc, eq_suff, 5.5)
    paragraph(doc,
        "자유형 오목은 판과 차례로 다음 합법 착수를 정할 수 있다. 단순 패 바둑은 직전 판 하나를 더 알아야 즉시 반복을 금지할 수 있다. positional superko처럼 "
        "더 긴 반복을 금지하는 규칙을 선택하면 residue도 그 규칙에 맞게 커져야 한다. 최소기억의 크기는 취향이 아니라 Law가 정한다.")

    heading(doc, "해석 탐색 예측의 구분", 2)
    table(doc,
        ["용어", "이 연구에서의 뜻", "쓰지 않는 뜻"],
        [
            ["상태전이", "합법행동을 적용해 다음 상태를 계산", "현실의 미래를 미리 관측"],
            ["후속 분기", "현재 상태에서 규칙상 생성 가능한 후보", "실제로 반드시 일어날 사건"],
            ["탐색", "후보 상태를 목적값으로 비교", "경험적 예측의 자동 성립"],
            ["residue 보존", "합법수 판정에 필요한 제약정보 유지", "미래 자체의 저장"],
            ["Renderer 평가", "현재 후보의 구조적 특징과 탐색 우선순위", "바둑 결과의 무조건적 예언"],
        ],
        [1.25, 2.7, 2.65], center_cols=(0,),
    )
    paragraph(doc,
        "따라서 이전 문서의 미래분기 보존이라는 표현은 사용하지 않는다. 정확한 문장은 다음과 같다. residue는 합법수 판정과 다음 상태의 전이 규칙을 정확히 보존한다. "
        "Renderer와 MCTS는 규칙이 허용한 후속 상태들을 계산해 행동을 고른다.")

    doc.add_page_break()
    heading(doc, "2 WRRA Game 0 1 구성적 분석", 1)
    heading(doc, "9×9 자유형 오목의 강제승리 출구", 2)
    o1 = r01["omok"]
    paragraph(doc,
        "0.1은 9×9 자유형 오목의 특정 합법 판에서 후보 71개를 모두 조사했다. 중앙 4 4 한 수만 다음 차례의 즉시승리 출구 네 개를 만들었다. 그 뒤 백의 "
        "가능한 응답 70개를 전부 적용했지만 네 출구를 동시에 없애는 응답은 없었다. 이 계산은 Renderer가 살아 있는 경계와 다중 위협을 이용해 후보를 줄이는 "
        "구성적 사례다.")
    figure(doc, ROOT / r01["figures"]["omok_fork"], "그림 1  9×9 자유형 오목에서 중앙 착수가 만든 네 개의 다음 수 승리 출구", 5.5)
    table(doc,
        ["항목", "값", "의미"],
        [
            ["합법 후보", o1["legal_move_count"], "모든 빈 칸"],
            ["다중출구 후보", o1["fork_candidate_count"], "중앙 4 4 한 수"],
            ["만든 승리 출구", o1["outlet_count"], "가로 세로 네 끝점"],
            ["검사한 방어", o1["defensive_replies_checked"], "중앙 뒤의 모든 백 응답"],
            ["모든 출구를 지운 방어", o1["defensive_replies_escaping_all_immediate_wins"], "강제승리 판정"],
            ["후보 감소율", f"{100*o1['candidate_reduction_fraction']:.2f}%", "구성적 사례의 Renderer 축약"],
        ],
        [2.0, 1.45, 3.15], center_cols=(1,),
    )

    heading(doc, "5×5 바둑의 패와 residue", 2)
    g1 = r01["go"]
    paragraph(doc,
        "바둑 사례는 현재 판의 모양만으로는 합법수를 정할 수 없다는 직접 반례를 구성했다. 한 수가 백돌 하나를 잡은 뒤 즉시 되따내면 직전 판이 복원된다. 직전 판 "
        "residue를 포함하면 합법행동은 18개이고, 이를 제거하면 금지되어야 할 되따내기가 들어와 19개가 된다.")
    figure(doc, ROOT / r01["figures"]["go_ko"], "그림 2  단순 패에서 직전 판 residue가 차단하는 즉시 되따내기", 6.1)
    table(doc,
        ["검사", "residue 포함", "residue 제거", "판정"],
        [["되따내기", "불법", "합법", "제거 시 직전 판 복원"],
         ["합법행동 수", g1["legal_actions_with_residue"], g1["legal_actions_without_residue"], "한 수 차이"]],
        [1.6, 1.55, 1.55, 2.05], center_cols=(0, 1, 2),
    )
    paragraph(doc,
        "0.1의 가치는 두 용어를 실제 계산에 연결했다는 점이다. Renderer는 살아 있는 위협과 경계를 드러내 후보를 정렬하거나 줄인다. residue는 현재 판만으로 보이지 않는 "
        "법적 제약을 상태에 결합한다. 0.2는 이 구성적 사례가 전체 도달상태에서도 유지되는지를 검사했다.")

    doc.add_page_break()
    heading(doc, "3 WRRA Game 0 2 완전탐색", 1)
    heading(doc, "4×4 3목 Connect K 실험계", 2)
    paragraph(doc,
        "표준 15×15 오목 전체를 열거하는 대신 4×4에서 연속 3개가 승리인 완전열거 실험계를 사용했다. 빈 판에서 시작해 턴 순서와 조기 종료를 지키며 도달 가능한 "
        "상태만 닫았다. 목적은 작은 판의 기력을 주장하는 것이 아니라 Renderer의 안전축약이 minimax 값을 바꾸는지 모든 상태에서 판정하는 것이다.")
    table(doc,
        ["장부", "값", "장부", "값"],
        [
            ["도달상태", f"{ck['reachable_full_states']:,}", "합법 방향전이", f"{ck['directed_legal_edges']:,}"],
            ["승리 종료", f"{ck['terminal_win_states']:,}", "무승부 종료", f"{ck['terminal_draw_states']:,}"],
            ["즉시승리 상태", f"{ck['phenotype_census']['immediate_win']:,}", "단일 강제방어", f"{ck['phenotype_census']['mandatory_block']:,}"],
            ["다중위협 강제패배", f"{ck['phenotype_census']['forced_loss_certificate']:,}", "열린 탐색", f"{ck['phenotype_census']['open_search']:,}"],
            ["전체 값 불일치", f"{ck['all_state_value_mismatches']:,}", "빈 판 값", ck["root_exact_value"]],
        ],
        [1.7, 1.3, 1.7, 1.3], center_cols=(1, 3),
    )
    equation(doc, eq_lossless, 5.75)
    paragraph(doc,
        "안전축약은 즉시승리 수가 있으면 그 수만 남기고, 상대의 즉시승리 지점이 하나면 그 지점만 막으며, 서로 다른 즉시승리 지점이 둘 이상이면 한 수로 모두 "
        "막을 수 없다는 증명서를 사용한다. 이 조건이 없는 상태에서는 수를 버리지 않고 순서만 바꾼다. 기준 minimax와 이 solver를 6,036,001개 상태 모두에서 "
        "비교한 결과 값 불일치는 0개였다.")
    figure(doc, figures02["search"], "그림 3  같은 빈 판 값을 얻기 위한 alpha beta 방문 노드 수", 6.15)
    table(doc,
        ["탐색", "방문 노드", "기준 대비 감소", "root 값"],
        [
            ["기준 행 우선", f"{ck['root_alpha_beta']['baseline']['nodes']:,}", "0%", ck["root_alpha_beta"]["baseline"]["value"]],
            ["Renderer 순서", f"{ck['root_alpha_beta']['renderer_order']['nodes']:,}", f"{100*ck['node_reduction_vs_baseline']['renderer_order_fraction']:.3f}%", ck["root_alpha_beta"]["renderer_order"]["value"]],
            ["Renderer 안전축약", f"{ck['root_alpha_beta']['renderer_gate']['nodes']:,}", f"{100*ck['node_reduction_vs_baseline']['renderer_gate_fraction']:.5f}%", ck["root_alpha_beta"]["renderer_gate"]["value"]],
        ],
        [2.0, 1.45, 1.8, 1.1], center_cols=(1, 2, 3),
    )

    heading(doc, "3×3 단순 패 바둑의 전체 도달그래프", 2)
    paragraph(doc,
        "포획, 자살수 금지, 단순 패, 패스, 두 번의 연속 패스 종료를 적용해 초기 빈 판에서 도달할 수 있는 완전상태를 모두 열거했다. 완전상태는 현재 판, 차례, "
        "직전 판 residue, 연속 패스 수다. 현재 판, 차례, 패스 수가 같은 상태를 묶은 뒤 residue만 달라질 때 합법 착수 집합이 달라지는지 비교했다.")
    figure(doc, figures02["go_counts"], "그림 4  같은 가시적 상태 안에서 residue와 합법수 집합이 달라지는 동치류", 6.1)
    table(doc,
        ["분류", "개수", "의미"],
        [
            ["완전상태", f"{go['reachable_full_states']:,}", "residue를 포함한 도달상태"],
            ["합법 방향전이", f"{go['directed_legal_edges']:,}", "착수와 패스"],
            ["가시적 상태 동치류", f"{go['unique_visible_state_classes']:,}", "판 차례 패스 수 동일"],
            ["복수 residue 동치류", f"{go['classes_with_multiple_reachable_residues']:,}", "서로 다른 직전 판"],
            ["합법수 차이 동치류", f"{go['classes_with_residue_dependent_legal_actions']:,}", "residue만으로 합법수 집합 변화"],
            ["패 차단 완전상태", f"{go['ko_sensitive_full_states']:,}", "각 상태에서 정확히 한 착수 차단"],
            ["최대 합법수 집합 수", f"{go['max_distinct_legal_action_sets_per_visible_state']:,}", "한 가시적 상태에 대응"],
        ],
        [2.2, 1.35, 3.0], center_cols=(1,),
    )
    figure(doc, figures02["witness"], "그림 5  같은 현재 판에서 residue에 따라 백의 0 1 합법성이 달라지는 최단 5수 증명쌍", 6.4)
    paragraph(doc,
        "이 결과는 예측 성능에 관한 것이 아니다. 같은 현재 판에서 어떤 수가 규칙상 허용되는지를 정확히 판정하려면 무엇을 상태에 포함해야 하는지에 관한 결과다. "
        "가장 짧은 반례의 두 이력은 모두 5수 뒤 같은 판과 같은 차례에 도달하지만, 한쪽의 백 0 1은 합법이고 다른 쪽에서는 직전 판을 복원하므로 금지된다.")

    heading(doc, "완전탐색이 만든 설계 규칙", 2)
    table(doc,
        ["제거시험", "관측된 변화", "설계 판정"],
        [
            ["Renderer 순서 제거", f"{ck['root_alpha_beta']['renderer_order']['nodes']:,} → {ck['root_alpha_beta']['baseline']['nodes']:,}노드", "정확도와 계산비용을 분리"],
            ["안전축약 제거", f"{ck['root_alpha_beta']['renderer_gate']['nodes']:,} → {ck['root_alpha_beta']['baseline']['nodes']:,}노드", "증명된 phenotype은 후보축약 가능"],
            ["직전 판 residue 제거", "768개 상태에서 금지수 1개 유입", "합법수 판정과 상태전이에 필수"],
            ["전체 이력 대신 직전 판", "단순 패 판정 유지", "Law가 요구한 최소기억"],
        ],
        [2.0, 2.2, 2.35], center_cols=(0,),
    )

    doc.add_page_break()
    heading(doc, "4 WRRA Go 0 1 실제 대국 알고리즘", 1)
    heading(doc, "완전탐색에서 실제 착수로", 2)
    paragraph(doc,
        "3×3 전체 열거는 상태정의와 규칙 구현을 검증하지만 9×9 이상에서는 전체 상태공간을 닫을 수 없다. WRRA Go 0.1은 검증된 상태전이를 그대로 사용하고, Renderer가 "
        "각 합법수의 구조적 특징을 수치화한 뒤 MCTS가 제한된 계산예산 안에서 후보를 반복 평가하도록 만들었다. 결과적으로 같은 프로그램이 5×5 시험대국과 9×9 "
        "사람 대 엔진 대국을 실행한다.")
    figure(doc, engine_figures["architecture"], "그림 6  WRRA Go 0.1의 상태에서 착수와 Ledger까지", 6.55)

    heading(doc, "완전상태와 합법전이", 2)
    equation(doc, eq_state, 4.8)
    table(doc,
        ["기호", "구성", "필요한 이유"],
        [
            ["B", "현재 흑백 점유 배열", "돌무리 활로 포획 득점"],
            ["p", "둘 차례", "행동 주체와 상대를 구분"],
            ["ρ", "직전 판 배열", "단순 패의 즉시 반복 금지"],
            ["c", "연속 패스 수", "두 패스 종료 판정"],
            ["n", "수 번호", "실행 장부와 운영상 탐색 경계"],
        ],
        [0.8, 2.6, 3.2], center_cols=(0,),
    )
    equation(doc, eq_update, 5.8)
    paragraph(doc,
        "착수 함수는 빈 칸인지 확인하고, 인접한 상대 돌무리의 활로가 0이면 포획하며, 새로 놓인 자기 돌무리의 활로가 0이면 자살수로 거부한다. 결과 판이 ρ와 같으면 "
        "단순 패로 거부한다. 패스는 판을 바꾸지 않되 직전 판과 연속 패스 수를 갱신한다. 두 번의 연속 패스에서 area score를 계산한다.")

    heading(doc, "투명한 Renderer 채널", 2)
    paragraph(doc,
        "Renderer는 신경망의 숨은 파라미터 대신 열한 개의 명시적 채널을 사용한다. 가중치는 학습하지 않고 고정했으며, 각 수의 점수와 phenotype을 Ledger에 남긴다. "
        "이 점수는 합법성을 정하지 않는다. 합법수 함수가 먼저 후보를 닫은 뒤 탐색 prior와 rollout 순서에만 사용한다.")
    equation(doc, eq_renderer, 3.8)
    weights = __import__("wrra_go_0_1").WRRARenderer.WEIGHTS
    feature_rows = [
        ["capture", "즉시 잡은 상대 돌 수", f"{weights['capture']:+.2f}"],
        ["rescue", "단수였던 인접 자기 돌의 구조적 구제", f"{weights['rescue']:+.2f}"],
        ["atari", "착수 뒤 단수가 된 상대 돌 수", f"{weights['atari']:+.2f}"],
        ["connection", "둘 이상의 자기 돌무리 연결", f"{weights['connection']:+.2f}"],
        ["cut pressure", "둘 이상의 상대 돌무리에 동시에 접촉", f"{weights['cut_pressure']:+.2f}"],
        ["liberties", "새 돌무리의 활로  최대 6", f"{weights['liberties']:+.2f}"],
        ["self atari", "포획 없는 한 집 활로", f"{weights['self_atari']:+.2f}"],
        ["eye fill", "포획 없이 자기 눈 후보를 메움", f"{weights['eye_fill']:+.2f}"],
        ["area delta", "작은 폐쇄영역의 즉시 변화  열린 대영역 제외", f"{weights['area_delta']:+.2f}"],
        ["position", "판 크기에 따른 초기 위치 선호", f"{weights['position']:+.2f}"],
        ["pass readiness", "상대 패스와 현재 우세를 이용한 종료", f"{weights['pass_readiness']:+.2f}"],
    ]
    table(doc, ["채널", "계산 내용", "가중치"], feature_rows, [1.45, 4.25, 0.9], center_cols=(0, 2))

    heading(doc, "MCTS 결합", 2)
    paragraph(doc,
        "MCTS 자체는 알려진 탐색 알고리즘이다. WRRA의 역할은 이를 새 이름으로 바꾸는 것이 아니라 충분한 상태, 합법전이, Renderer prior, guided rollout, phenotype과 "
        "Ledger를 한 실행계약 안에서 제공하는 데 있다. 같은 simulation 수의 균등 MCTS를 통제군으로 두어 Renderer가 추가한 차이를 직접 비교할 수 있다.")
    equation(doc, eq_puct, 6.15)
    numbered(doc, [
        "현재 완전상태에서 착수와 패스를 포함한 합법전이를 생성한다.",
        "Renderer 점수를 softmax로 정규화해 후보 prior를 만든다.",
        "방문가치 Q와 탐색항을 합친 기준으로 트리를 내려간다.",
        "새 노드에서는 Renderer 상위 후보를 중심으로 제한된 rollout을 실행한다.",
        "종료 또는 계산 경계의 area score를 root 플레이어 관점의 값으로 되돌린다.",
        "가장 많이 방문한 수를 선택하고 상위 후보의 방문 수 가치 prior phenotype을 기록한다.",
    ])

    heading(doc, "실제로 사용하는 방법", 2)
    table(doc,
        ["목적", "명령", "결과"],
        [
            ["9×9 사람 대 엔진", "python wrra_go_0_1.py play --size 9 --human B", "좌표 입력형 대국"],
            ["규칙 회귀검사", "python wrra_go_0_1.py test", "패 포획 자살수 종료 검사"],
            ["색상 교대 벤치마크", "python wrra_go_0_1.py benchmark", "JSON 성능 장부"],
            ["9×9 자기대국", "python wrra_go_0_1.py selfplay --size 9", "JSON과 SGF 대국기보"],
        ],
        [1.55, 3.8, 1.65], center_cols=(0,),
    )
    paragraph(doc,
        "사람은 D4 같은 바둑 좌표와 pass를 입력한다. I열은 통상적인 바둑 좌표처럼 건너뛴다. 기본 72 simulations는 실행가능성을 위한 설정이며 더 높은 수를 주면 한 수당 "
        "계산시간을 늘려 더 많은 후보를 검사한다. 13×13과 19×19도 규칙상 실행할 수 있지만 이번 버전의 정량시험은 5×5와 9×9에 집중했다.")

    heading(doc, "5 실행검증", 1)
    heading(doc, "결정적 규칙검사", 2)
    test_labels = {
        "same_visible_board_witness": "같은 현재 판 증명쌍 재현",
        "simple_ko_blocks_repetition": "단순 패가 직전 판 복원을 차단",
        "removing_residue_admits_repetition": "residue 제거 시 반복수가 유입",
        "same_board_different_residue_changes_legality": "같은 판에서 residue가 합법성 변경",
        "two_passes_terminate": "두 번의 연속 패스 종료",
        "capture_removes_surrounded_group": "포획 시 돌무리 제거",
        "suicide_is_forbidden": "자살수 금지",
        "area_scoring_is_deterministic": "area score 결정성",
        "renderer_does_not_claim_open_board": "초반 열린 빈 영역의 영토 오인 방지",
        "coordinate_round_trip": "좌표 변환 왕복",
    }
    table(doc,
        ["검사항목", "결과"],
        [[test_labels[key], "PASS" if value else "FAIL"] for key, value in tests.items()],
        [5.4, 1.2], center_cols=(1,),
    )
    paragraph(doc,
        "이 가운데 핵심 회귀검사는 Game 0.2가 찾은 최단 5수 반례를 새 엔진의 독립된 상태객체로 다시 재생하는 것이다. 두 이력은 같은 판과 같은 차례에 도달하지만 "
        "직전 판이 다르다. 동일한 백 0 1 착수는 한 상태에서 합법이고 다른 상태에서 패로 거부된다. 전수검증에서 얻은 상태조건이 실제 대국 코드에 이어졌음을 확인한다.")

    heading(doc, "5×5 색상 교대 벤치마크", 2)
    paragraph(doc,
        f"WRRA MCTS와 통제 알고리즘을 흑백으로 번갈아 배치했다. 각 수는 {benchmark['search']['simulations_per_move']} simulations를 사용했다. 균등 MCTS는 "
        "같은 트리탐색과 같은 계산예산을 사용하지만 Renderer prior와 guided rollout을 제거했다. 무작위 통제군은 합법 착수 가운데 하나를 고른다.")
    figure(doc, engine_figures["benchmark"], "그림 7  동일 규칙과 색상 교대 조건에서의 WRRA MCTS 대조 대국", 5.7)
    table(doc,
        ["대조군", "대국", "WRRA 승패", "승률", "Wilson 95% 구간", "평균 수수"],
        [
            ["균등 MCTS", uniform["games"], f"{uniform['wrra_wins']}승 {uniform['wrra_losses']}패", f"{100*uniform['wrra_win_rate']:.1f}%", f"{100*uniform['wilson_95_interval'][0]:.1f}–{100*uniform['wilson_95_interval'][1]:.1f}%", f"{uniform['mean_move_count']:.1f}"],
            ["무작위", random_summary["games"], f"{random_summary['wrra_wins']}승 {random_summary['wrra_losses']}패", f"{100*random_summary['wrra_win_rate']:.1f}%", f"{100*random_summary['wilson_95_interval'][0]:.1f}–{100*random_summary['wilson_95_interval'][1]:.1f}%", f"{random_summary['mean_move_count']:.1f}"],
        ],
        [1.2, 0.75, 1.25, 0.85, 1.65, 1.0], center_cols=(0, 1, 2, 3, 4, 5),
    )
    paragraph(doc,
        "이 시험으로 확정되는 것은 고정 Renderer가 실제 수 선택에 유효한 편향을 주었고, 같은 계산예산의 균등 탐색보다 이 표본에서 더 많은 대국을 이겼다는 사실이다. "
        "전문 엔진보다 강하다는 결론이나 19×19 기력의 일반화는 이 표본에서 꺼내지 않는다. 다음 버전은 판 크기와 seed를 늘리고 GTP 기반 외부 엔진 대국을 추가하면 된다.")

    heading(doc, "9×9 자기대국", 2)
    paragraph(doc,
        f"WRRA MCTS 두 개가 각각 {selfplay['moves'][0]['decision']['simulations']} simulations로 둔 9×9 자기대국은 {selfplay['move_count']}수 뒤 "
        f"{selfplay['ended_by']}로 끝났다. 흑 면적 {selfplay['score']['black_area']}, 백 면적 {selfplay['score']['white_area']}, 덤 {selfplay['score']['komi']}를 적용한 "
        f"흑 기준 결과는 {selfplay['score']['black_margin']:+.1f}집이며 승자는 {selfplay['winner']}이다.")
    figure(doc, engine_figures["final_board"], "그림 8  9×9 WRRA 자기대국의 최종 판  SGF 기보를 별도 제공한다", 4.9)
    figure(doc, engine_figures["candidates"], "그림 9  첫 착수에서 실제로 남긴 후보별 방문 장부", 5.9)
    first_candidates = selfplay["moves"][0]["decision"]["candidates"][:8]
    table(doc,
        ["후보", "방문", "평균값", "prior", "Renderer 점수", "phenotype"],
        [[item["action"], item["visits"], f"{item['value']:+.3f}", f"{item['prior']:.4f}", f"{item['renderer_score']:+.3f}", item["phenotype"]] for item in first_candidates],
        [0.8, 0.75, 1.0, 1.0, 1.35, 1.35], center_cols=(0, 1, 2, 3, 4, 5),
    )

    heading(doc, "전체 자기대국 기보", 2)
    moves = selfplay["moves"]
    paired_moves = []
    for i in range(0, len(moves), 2):
        black_move = moves[i]["action"] if i < len(moves) else ""
        white_move = moves[i + 1]["action"] if i + 1 < len(moves) else ""
        paired_moves.append([i // 2 + 1, black_move, white_move])
    table(doc, ["수순", "흑", "백"], paired_moves, [1.0, 2.8, 2.8], center_cols=(0, 1, 2))
    paragraph(doc,
        "SGF는 같은 수순을 표준 좌표로 저장한다. JSON Ledger에는 각 수의 상위 후보, 방문 수, 평균값, prior, Renderer 점수, phenotype과 계산시간이 포함된다. "
        "따라서 최종 착수뿐 아니라 왜 그 수가 선택되었는지 다시 검사할 수 있다.")

    doc.add_page_break()
    heading(doc, "6 이 연구의 학문적 가치", 1)
    heading(doc, "하나의 추상어가 아니라 연결된 검증 사슬", 2)
    numbered(doc, [
        "같은 Core 계약을 서로 다른 게임에 적용하되 필요한 residue의 차이를 보존했다. 공통성은 차이를 지우는 통합이 아니라 같은 검사법을 공유하는 통합이다.",
        "구성적 사례에서 시작해 전체 도달그래프의 반례와 값 보존으로 확장했다. 설명이 계산결과에 의해 수정될 수 있는 구조를 갖췄다.",
        "상태의 정확성과 탐색의 효율성을 분리했다. residue 제거는 합법성을 바꾸고 Renderer 제거는 값은 유지한 채 계산량을 늘린다.",
        "완전탐색이 불가능한 판에서는 검증된 전이기를 버리지 않고 근사탐색과 결합했다. 작은 판의 exact proof와 큰 판의 bounded action을 연결했다.",
        "Renderer의 채널과 가중치, 후보별 방문 수를 공개해 숨은 판단을 줄였다. 각 착수의 구조적 근거를 phenotype으로 되짚을 수 있다.",
        "오목 바둑 장기 체스 같은 서로 다른 게임을 state law residue renderer ledger라는 동일 인터페이스에 올릴 수 있는 재사용 가능한 연구도구를 만들었다.",
    ])

    heading(doc, "새롭게 드러난 장점", 2)
    table(doc,
        ["장점", "이번 연구의 증거", "다른 복잡계로의 의미"],
        [
            ["상태 aliasing 자동탐지", "752개 residue 의존 동치류", "같은 관측 뒤 다른 허용행동 탐색"],
            ["증명형 후보축약", "Connect K 전 상태 값 불일치 0", "국소 phenotype을 안전한 계산인증서로 사용"],
            ["정확성과 속도의 분리", "residue와 Renderer 제거시험", "모델 기억과 계산예산을 별도 설계"],
            ["축척 전환", "3×3 exact에서 9×9 MCTS로 이동", "전수검증과 제한계산의 연결"],
            ["설명가능한 행동", "11개 채널과 후보 Ledger", "결과뿐 아니라 선택근거 재검사"],
            ["Core 불변성", "0.1 0.2 Go 0.1 모두 Core 수정 0", "응용이 늘어도 상위 계약 유지"],
        ],
        [1.55, 2.25, 2.45], center_cols=(0,),
    )
    paragraph(doc,
        "특히 작은 완전계에서 상태표현과 축약규칙을 검증한 뒤 같은 코드를 실제 행위자에 넣는 순서는 중요하다. 기력이 먼저가 아니라 규칙의 정확성, 상태의 충분성, "
        "축약의 무손실성, 제한계산의 성능을 차례로 분리한다. 이 순서는 전력망, 생물학적 조절, 에이전트 행동처럼 전체공간을 열거할 수 없는 영역에서도 작은 검증계와 "
        "실제 운영계를 연결하는 방법으로 확장할 수 있다.")

    heading(doc, "현재 성과와 다음 검증을 구분하는 기준", 2)
    table(doc,
        ["현재 확정", "다음에 확장"],
        [
            ["단순 패를 포함한 합법수와 상태전이", "positional superko의 압축 residue"],
            ["5×5 통제 대국과 9×9 실제 실행", "9×9 다중 seed 대회와 19×19 시간제한 대국"],
            ["고정된 투명 Renderer 채널", "가중치 ablation과 자동보정"],
            ["MCTS 결합의 실행 가능성", "KataGo 등 외부 GTP 엔진과 등급화"],
            ["area score와 SGF JSON 장부", "사석 판정과 중국식 일본식 규칙 profile 분리"],
        ],
        [3.3, 3.3], center_cols=(),
    )
    paragraph(doc,
        "이 구분은 현재 성과를 낮추기 위한 문장이 아니라 다음 실험의 입력과 출력이 무엇인지 고정하는 장부다. 현재 버전은 이미 사람이 둘 수 있는 독립 프로그램이며, "
        "전문 대국 성능은 같은 실행계약 위에서 계산예산과 평가함수를 확장해 측정할 후속 항목이다.")

    doc.add_page_break()
    heading(doc, "7 재현과 Zenodo 공개 구성", 1)
    heading(doc, "파일 계보", 2)
    manifest = [
        ["wrra_game_0_1.py", "구성적 오목과 바둑 분석", file_hash(ROOT / "wrra_game_0_1.py")],
        ["wrra_game_0_1_results.json", "0.1 결과 장부", file_hash(R01_PATH)],
        ["wrra_game_0_2.py", "Connect K와 3×3 바둑 전수검사", file_hash(ROOT / "wrra_game_0_2.py")],
        ["wrra_game_0_2_results.json", "0.2 결과 장부", file_hash(R02_PATH)],
        ["wrra_go_0_1.py", "실제 대국 엔진과 통제시험", file_hash(ROOT / "wrra_go_0_1.py")],
        [BENCH_PATH.name, "색상 교대 벤치마크", file_hash(BENCH_PATH)],
        [TEST_PATH.name, "결정적 규칙 회귀검사", file_hash(TEST_PATH)],
        [SELFPLAY_PATH.name, "9×9 자기대국 전체 Ledger", file_hash(SELFPLAY_PATH)],
        [SELFPLAY_PATH.with_suffix(".sgf").name, "9×9 자기대국 SGF", file_hash(SELFPLAY_PATH.with_suffix(".sgf"))],
    ]
    table(doc, ["파일", "역할", "SHA 256"], manifest, [2.55, 2.05, 2.25], center_cols=(0,))

    heading(doc, "재현 순서", 2)
    numbered(doc, [
        "python wrra_go_0_1.py test를 실행해 열 개 규칙검사가 모두 true인지 확인한다.",
        "python wrra_game_0_2.py를 실행해 Connect K 6,036,001상태와 값 불일치 0을 확인한다.",
        "같은 결과에서 3×3 바둑 132,161상태와 residue 의존 동치류 752개를 확인한다.",
        "python wrra_go_0_1.py benchmark --size 5 --paired-games 8 --simulations 48을 실행한다.",
        "python wrra_go_0_1.py selfplay --size 9 --simulations 24를 실행해 JSON과 SGF를 만든다.",
        "사람 대 엔진은 python wrra_go_0_1.py play --size 9 --human B로 시작한다.",
    ])
    paragraph(doc,
        f"이번 벤치마크의 기록된 실행시간은 {benchmark['elapsed_seconds']:.2f}초이며 환경에 따라 달라진다. 상태 수, 합법성 검사, seed와 대국 결과는 같은 Python 동작과 "
        "입력 조건에서 재현된다. 0.2의 완전탐색 결과는 기존 JSON을 포함했으므로 전체 재계산 없이도 장부를 검사할 수 있다.")

    heading(doc, "Zenodo 권장 메타데이터", 2)
    table(doc,
        ["필드", "권장 입력"],
        [
            ["Title", "WRRA Game 1.0 Minimal Sufficient State Exhaustive Validation and a Playable Go Engine"],
            ["Subtitle", "From Omok and Go State Analysis to WRRA Go 0.1"],
            ["Creator", "Wonsik Choi"],
            ["Contact", "janefather@gmail.com"],
            ["Resource type", "Software and technical report"],
            ["Language", "Korean with English metadata"],
            ["Keywords", "WRRA, Go, Baduk, Omok, Connect K, MCTS, state sufficiency, residue, renderer, exhaustive search, complex systems"],
            ["Related identifier", "WRRA Core 1.0  https://doi.org/10.5281/zenodo.22650956"],
            ["License", "보고서·문서·데이터: CC BY 4.0; 소스 코드: MIT"],
        ],
        [1.45, 5.15], center_cols=(0,),
    )
    heading(doc, "Description", 3)
    paragraph(doc,
        "This release integrates the WRRA Game research program from constructive game analysis to exhaustive finite-state validation and a playable Go engine. "
        "WRRA Game 0.1 identifies an Omok forced-win phenotype and a simple-ko residue counterexample. WRRA Game 0.2 exhaustively enumerates 6,036,001 reachable "
        "Connect-K states and 132,161 reachable 3×3 Go states, finding zero minimax-value mismatches for the certified Renderer reduction and 752 visible-state classes "
        "whose legal actions depend on the previous-board residue. WRRA Go 0.1 reuses the verified state transition system in a playable 9×9 engine, combining a transparent "
        "eleven-channel Renderer with Monte Carlo Tree Search. The archive includes source code, machine-readable ledgers, deterministic rule tests, paired-control matches, and an SGF "
        "self-play record. The work evaluates rule-permitted continuations; it does not describe the residue or search as prediction of a pre-existing future.")

    heading(doc, "참조 계보", 2)
    paragraph(doc, "WRRA Core 1.0  Frozen Domain Agnostic Execution Architecture  DOI 10.5281/zenodo.22650956")
    paragraph(doc, "WRRA Game 0.1  오목과 바둑에서 단순 규칙이 만드는 복잡성  2026년 9월 16일")
    paragraph(doc, "WRRA Game 0.2  오목형 연결게임 완전탐색과 바둑 residue 상태충분성 검증  2026년 9월 17일")
    paragraph(doc, "WRRA Go 0.1  본 보고서에 포함된 실제 대국 엔진  2026년 9월 17일")

    heading(doc, "최종 판정", 1)
    paragraph(doc,
        "이 연구는 단순규칙이 복잡한 게임을 만든다는 설명에서 멈추지 않았다. 오목에서는 Renderer가 어떤 phenotype을 안전한 계산축약으로 바꿀 수 있는지 전체 상태에서 "
        "검사했고, 바둑에서는 어떤 residue가 합법수 판정을 위해 반드시 상태에 들어가야 하는지 전체 도달그래프의 반례로 확정했다. 이어 같은 전이기를 9×9 대국 엔진에 "
        "넣어 실제 수를 선택하고 SGF와 후보 Ledger를 남겼다.")
    paragraph(doc,
        "따라서 WRRA Game 1.0의 핵심 성과는 한 문장으로 정리된다. 최소 충분상태로 규칙의 정확성을 지키고, Renderer로 계산해야 할 구조를 드러내며, 제한된 탐색으로 "
        "실제 행동을 선택하는 하나의 재현 가능한 게임 실행계를 구축했다.")

    props = doc.core_properties
    props.title = "WRRA Game 1 0 복잡게임의 최소 충분상태와 실제 바둑 엔진"
    props.subject = "Integrated WRRA Game research and playable Go engine"
    props.author = "Wonsik Choi"
    props.keywords = "WRRA, Go, Baduk, Omok, Connect K, MCTS, state sufficiency, residue, renderer, exhaustive search"
    props.comments = "WRRA Core 1.0 frozen; integrated Zenodo release"
    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build()
